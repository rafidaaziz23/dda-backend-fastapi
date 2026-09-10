import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load datasets
files = [
    ('OFF', 'prototype/garden-rampage/dataset/telemetry_session_2026-09-11_00-00-48.json'),
    ('FCM', 'prototype/garden-rampage/dataset/telemetry_session_2026-09-11_00-10-46.json'),
    ('GMM', 'prototype/garden-rampage/dataset/telemetry_session_2026-09-11_00-21-28.json')
]

keys = [
    'avg_hp_remaining_pct', 'total_kills', 'accuracy_pct', 'dash_frequency',
    'avg_time_per_wave_sec', 'damage_taken_total', 'near_death_events',
    'hits_taken_from_strawberry', 'hits_taken_from_jambu', 'avg_enemies_alive_simultaneously'
]

col_headers = [
    'No', 'Skenario', 'Wave',
    'Avg HP Remaining (%)', 'Total Kills', 'Accuracy (%)', 'Dash Freq (/min)',
    'Avg Time (s)', 'Damage Taken', 'Near Death Events',
    'Hits Strawberry', 'Hits Jambu', 'Avg Enemies Alive'
]

wb = openpyxl.Workbook()

# Define styles
header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')

sub_header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
off_fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
fcm_fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
gmm_fill = PatternFill(start_color='ED7D31', end_color='ED7D31', fill_type='solid')

data_font = Font(name='Arial', size=10)
bold_data_font = Font(name='Arial', size=10, bold=True)
center_align = Alignment(horizontal='center', vertical='center')
left_align = Alignment(horizontal='left', vertical='center')
right_align = Alignment(horizontal='right', vertical='center')

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# -------------------------------------------------------------
# Sheet 1: Master Dataset (63 Wave)
# -------------------------------------------------------------
ws1 = wb.active
ws1.title = "Master Dataset (63 Wave)"

# Title banner
ws1.merge_cells('A1:M1')
ws1['A1'] = "DATASET PENGUJIAN KOMPARATIF DDA (GARDEN RAMPAGE - 63 WAVE)"
ws1['A1'].font = Font(name='Arial', size=14, bold=True, color='1F4E79')
ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[1].height = 30

ws1.append([])

# Header row
ws1.append(col_headers)
ws1.row_dimensions[3].height = 24
for col_idx in range(1, len(col_headers) + 1):
    cell = ws1.cell(row=3, column=col_idx)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align

row_num = 1
current_excel_row = 4

for label, path in files:
    with open(path) as f:
        data = json.load(f)
    for d in data:
        w = d['wave']
        row_data = [
            row_num,
            label,
            w,
            d['avg_hp_remaining_pct'],
            d['total_kills'],
            d['accuracy_pct'],
            d['dash_frequency'],
            d['avg_time_per_wave_sec'],
            d['damage_taken_total'],
            d['near_death_events'],
            d['hits_taken_from_strawberry'],
            d['hits_taken_from_jambu'],
            d['avg_enemies_alive_simultaneously']
        ]
        ws1.append(row_data)
        
        # Style row
        ws1.row_dimensions[current_excel_row].height = 20
        for c in range(1, len(row_data) + 1):
            cell = ws1.cell(row=current_excel_row, column=c)
            cell.font = data_font
            cell.border = thin_border
            if c in [1, 3]:
                cell.alignment = center_align
            elif c == 2:
                cell.alignment = center_align
                if label == 'OFF':
                    cell.font = Font(name='Arial', size=10, bold=True, color='9C0006')
                elif label == 'FCM':
                    cell.font = Font(name='Arial', size=10, bold=True, color='002060')
                else:
                    cell.font = Font(name='Arial', size=10, bold=True, color='C65911')
            elif c in [4, 6]:
                cell.number_format = '0.00%'
                cell.alignment = right_align
            elif c in [7, 8, 13]:
                cell.number_format = '0.00'
                cell.alignment = right_align
            else:
                cell.number_format = '#,##0'
                cell.alignment = right_align
                
        row_num += 1
        current_excel_row += 1

# Auto-fit columns
for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

# -------------------------------------------------------------
# Sheet 2: Ringkasan Statistik
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="Ringkasan Statistik")

ws2.merge_cells('A1:D1')
ws2['A1'] = "RINGKASAN STATISTIK DESKRIPTIF PERBANDINGAN ALGORITMA DDA"
ws2['A1'].font = Font(name='Arial', size=13, bold=True, color='1F4E79')
ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 28

ws2.append([])

stat_headers = ['Parameter Evaluasi', 'OFF (Kontrol)', 'FCM (Usulan)', 'GMM (Pembanding)']
ws2.append(stat_headers)
ws2.row_dimensions[3].height = 24
for c in range(1, 5):
    cell = ws2.cell(row=3, column=c)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align

stats_rows = [
    ("Total Gelombang (Wave)", "27 Wave", "18 Wave", "18 Wave"),
    ("Titik Gugur (Death Wave)", "Wave 27 (Spike 3)", "Wave 18 (Spike 2)", "Wave 18 (Spike 2)"),
    ("Rata-rata Sisa HP (Mean)", "61.83%", "67.63% (Tertinggi)", "63.51%"),
    ("Variansi HP (Std Dev)", "0.2228", "0.2016 (Paling Stabil)", "0.2576 (Paling Fluktuatif)"),
    ("Retensi Flow Zone (40% - 80%)", "59.26% (16 wave)", "66.67% (12 wave - Tertinggi)", "33.33% (6 wave - Terpolarisasi)"),
    ("Kondisi Struggling (HP < 40%)", "14.81%", "11.11%", "22.22%"),
    ("Kondisi Dominant (HP > 80%)", "25.93%", "22.22%", "44.44%"),
    ("Total Near-Death Events", "6 kali", "1 kali (Paling Aman)", "4 kali"),
    ("Rata-rata Damage / Wave", "22.81 poin", "18.11 poin", "17.22 poin"),
    ("Rata-rata Akurasi Tembakan", "91.57%", "92.86%", "94.39%"),
    ("Frekuensi Dash Rata-rata", "8.96 kali/menit", "15.22 kali/menit (Aktif)", "10.48 kali/menit"),
    ("Rata-rata Durasi per Wave", "32.18 detik", "26.37 detik", "24.91 detik"),
    ("Silhouette Score (Baseline)", "—", "0.3644", "0.4300"),
    ("Davies-Bouldin Index (Baseline)", "—", "0.9412", "0.8443")
]

for idx, r in enumerate(stats_rows):
    ws2.append(list(r))
    curr_r = idx + 4
    ws2.row_dimensions[curr_r].height = 20
    for c in range(1, 5):
        cell = ws2.cell(row=curr_r, column=c)
        cell.font = data_font
        cell.border = thin_border
        if c == 1:
            cell.font = bold_data_font
            cell.alignment = left_align
        else:
            cell.alignment = center_align
            if c == 3 and "Tertinggi" in str(cell.value):
                cell.font = Font(name='Arial', size=10, bold=True, color='006100')
                cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

for col in ws2.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = max(max_len + 4, 18)

# -------------------------------------------------------------
# Sheet 3: Tabel 4 Min-Max (Format Mahmud)
# -------------------------------------------------------------
ws3 = wb.create_sheet(title="Tabel 4 Min-Max (Format Mahmud)")

ws3.merge_cells('A1:F1')
ws3['A1'] = "TABEL 4. TRANSFORMASI DATA TELEMETRI MENGGUNAKAN NORMALISASI MIN-MAX (SAMPEL WAVE 1)"
ws3['A1'].font = Font(name='Arial', size=12, bold=True, color='1F4E79')
ws3['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 26

ws3.append([])

t4_headers = ['No', 'Fitur Telemetri', 'Nilai Mentah (x) [Sebelum]', 'Batas Min (x_min)', 'Batas Max (x_max)', 'Nilai Normalisasi (x_norm) [Sesudah]']
ws3.append(t4_headers)
ws3.row_dimensions[3].height = 24
for c in range(1, 7):
    cell = ws3.cell(row=3, column=c)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align

t4_data = [
    (1, "avg_hp_remaining_pct", 1.00, 0.16, 1.00, 1.0000),
    (2, "total_kills", 5, 5, 57, 0.0000),
    (3, "accuracy_pct", 0.84, 0.77, 1.00, 0.3158),
    (4, "dash_frequency", 0.00, 0.00, 29.47, 0.0000),
    (5, "avg_time_per_wave_sec", 10.06, 8.20, 93.17, 0.0219),
    (6, "damage_taken_total", 0, 0, 82, 0.0000),
    (7, "near_death_events", 0, 0, 2, 0.0000),
    (8, "hits_taken_from_strawberry", 0, 0, 0, 0.0000),
    (9, "hits_taken_from_jambu", 0, 0, 1, 0.0000),
    (10, "avg_enemies_alive_simultaneously", 1.70, 0.88, 15.55, 0.0562)
]

for idx, r in enumerate(t4_data):
    ws3.append(list(r))
    curr_r = idx + 4
    ws3.row_dimensions[curr_r].height = 20
    for c in range(1, 7):
        cell = ws3.cell(row=curr_r, column=c)
        cell.font = data_font
        cell.border = thin_border
        if c == 1:
            cell.alignment = center_align
        elif c == 2:
            cell.alignment = left_align
            cell.font = Font(name='Consolas', size=10)
        elif c == 6:
            cell.font = bold_data_font
            cell.alignment = right_align
            cell.number_format = '0.0000'
            cell.fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
        else:
            cell.alignment = right_align
            if isinstance(cell.value, float):
                cell.number_format = '0.00'

for col in ws3.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws3.column_dimensions[col_letter].width = max(max_len + 4, 15)

# -------------------------------------------------------------
# Sheet 4: Tabel Ringkas (1 s.d. 10) - Opsi A Portrait
# -------------------------------------------------------------
ws4 = wb.create_sheet(title="Tabel Ringkas (1 s.d. 10)")

ws4.merge_cells('A1:M1')
ws4['A1'] = "TABEL DATA TELEMETRI LENGKAP 63 WAVE (FORMAT RINGKAS 1 s.d. 10)"
ws4['A1'].font = Font(name='Arial', size=12, bold=True, color='1F4E79')
ws4['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws4.row_dimensions[1].height = 26

ws4.append([])

headers_ringkas = ['No', 'Skenario', 'Wave', '1', '2', '3', '4', '5', '6', '7', '8', '9', '10']
ws4.append(headers_ringkas)
ws4.row_dimensions[3].height = 22
for c in range(1, len(headers_ringkas) + 1):
    cell = ws4.cell(row=3, column=c)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align

row_num = 1
current_r4 = 4

for label, path in files:
    with open(path) as f:
        data = json.load(f)
    for d in data:
        w = d['wave']
        r_data = [
            row_num,
            label,
            w,
            d['avg_hp_remaining_pct'],
            d['total_kills'],
            d['accuracy_pct'],
            d['dash_frequency'],
            d['avg_time_per_wave_sec'],
            d['damage_taken_total'],
            d['near_death_events'],
            d['hits_taken_from_strawberry'],
            d['hits_taken_from_jambu'],
            d['avg_enemies_alive_simultaneously']
        ]
        ws4.append(r_data)
        ws4.row_dimensions[current_r4].height = 19
        for c in range(1, len(r_data) + 1):
            cell = ws4.cell(row=current_r4, column=c)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align
            if c == 2:
                if label == 'OFF':
                    cell.font = Font(name='Arial', size=10, bold=True, color='9C0006')
                elif label == 'FCM':
                    cell.font = Font(name='Arial', size=10, bold=True, color='002060')
                else:
                    cell.font = Font(name='Arial', size=10, bold=True, color='C65911')
            elif c in [4, 6]:
                cell.number_format = '0.00'
            elif c in [7, 8, 13]:
                cell.number_format = '0.0'
        row_num += 1
        current_r4 += 1

# Keterangan Nomor Kolom di bawah tabel
current_r4 += 1
ws4.cell(row=current_r4, column=1, value="Keterangan Kolom (Fitur Telemetri):").font = bold_data_font
current_r4 += 1
legend_text = (
    "1: avg_hp_remaining_pct | 2: total_kills | 3: accuracy_pct | 4: dash_frequency (/menit) | "
    "5: avg_time_per_wave_sec (detik) | 6: damage_taken_total | 7: near_death_events | "
    "8: hits_taken_from_strawberry | 9: hits_taken_from_jambu | 10: avg_enemies_alive_simultaneously"
)
ws4.merge_cells(start_row=current_r4, start_column=1, end_row=current_r4, end_column=13)
ws4.cell(row=current_r4, column=1, value=legend_text).font = Font(name='Arial', size=9, italic=True)

# Lebar kolom ringkas (sempit dan hemat tempat)
ws4.column_dimensions['A'].width = 6   # No
ws4.column_dimensions['B'].width = 10  # Skenario
ws4.column_dimensions['C'].width = 8   # Wave
for col_char in ['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']:
    ws4.column_dimensions[col_char].width = 7.5

# Save to output locations
out_paths = [
    'D:/Colleges/TA/dda_fcm_gmm/DATASET_LENGKAP_63_WAVE.xlsx',
    'D:/Colleges/TA/dda_fcm_gmm/prototype/garden-rampage/dataset/DATASET_LENGKAP_63_WAVE.xlsx'
]

for p in out_paths:
    try:
        wb.save(p)
        print(f"[SUCCESS] File Excel tersimpan di: {p}")
    except PermissionError:
        print(f"[LOCKED] File {p} sedang dibuka di aplikasi Excel. Tutup Excel lalu jalankan lagi jika ingin menimpa file tersebut.")
